from io import BytesIO

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch


GREEN = "059669"
DARK_GREEN = "064E3B"
PALE_GREEN = "ECFDF5"
PALE_TEAL = "F0FDFA"
TEXT = "1F2937"


def _set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_run_font(run, size=11, color=TEXT, bold=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _make_diagram(labels, title, colors=None):
    colors = colors or [GREEN] * len(labels)
    figure = Figure(figsize=(10, 2.15), dpi=180, facecolor="white")
    canvas = FigureCanvasAgg(figure)
    axis = figure.add_axes([0.02, 0.08, 0.96, 0.78])
    axis.set_xlim(0, len(labels))
    axis.set_ylim(0, 1)
    axis.axis("off")
    figure.suptitle(title, fontsize=12, fontweight="bold", color=f"#{DARK_GREEN}")

    for index, label in enumerate(labels):
        x = index + 0.08
        box = FancyBboxPatch(
            (x, 0.25), 0.78, 0.48,
            boxstyle="round,pad=0.04,rounding_size=0.08",
            linewidth=1.4,
            edgecolor=f"#{colors[index]}",
            facecolor="#ECFDF5" if index % 2 == 0 else "#F0FDFA",
        )
        axis.add_patch(box)
        axis.text(x + 0.39, 0.49, label, ha="center", va="center",
                  fontsize=8.2, color=f"#{DARK_GREEN}", wrap=True)
        if index < len(labels) - 1:
            axis.add_patch(FancyArrowPatch(
                (x + 0.8, 0.49), (x + 0.98, 0.49),
                arrowstyle="-|>", mutation_scale=10,
                linewidth=1.2, color=f"#{GREEN}",
            ))

    output = BytesIO()
    canvas.print_png(output)
    output.seek(0)
    return output


def _make_risk_diagram():
    return _make_diagram(
        ["Xanh\nThấp", "Vàng\nThận trọng", "Cam\nNguy cơ cao", "Đỏ\nRất cao"],
        "Bốn mức cảnh báo rủi ro",
        ["059669", "D4A017", "EA7B19", "DC2626"],
    )


def _insert_illustrations(doc):
    paragraphs = doc.paragraphs
    placeholder_index = next(
        (i for i, paragraph in enumerate(paragraphs)
         if paragraph.text.startswith("[BẠN CÓ THỂ CHÈN ẢNH")),
        None,
    )
    if placeholder_index is None:
        raise ValueError("Không tìm thấy vị trí chèn hình minh họa trong báo cáo.")

    placeholder = paragraphs[placeholder_index]
    old_caption = paragraphs[placeholder_index + 1] if placeholder_index + 1 < len(paragraphs) else None
    anchor = paragraphs[placeholder_index - 1] if placeholder_index else placeholder
    illustrations = (
        (
            _make_diagram(
                ["Nhập địa chỉ", "Chuẩn hóa", "55 dấu hiệu", "4 mô hình", "Mức cảnh báo"],
                "Quy trình phân tích địa chỉ trên thiết bị người dùng",
            ),
            "Hình 1. Quy trình phân tích địa chỉ của SafeWeb.",
        ),
        (
            _make_diagram(
                ["Dữ liệu đã biết", "Đặc điểm URL", "Ý kiến mô hình", "Kết quả giải thích"],
                "Ba lớp thông tin hỗ trợ đánh giá rủi ro",
            ),
            "Hình 2. Các lớp thông tin được kết hợp trong quá trình đánh giá.",
        ),
        (
            _make_diagram(
                ["Trình duyệt", "Phân tích cục bộ", "Hiển thị kết quả"],
                "Luồng xử lý mặc định không gửi địa chỉ ra ngoài",
            ),
            "Hình 3. Luồng xử lý cục bộ nhằm bảo vệ quyền riêng tư.",
        ),
        (
            _make_risk_diagram(),
            "Hình 4. Bốn mức cảnh báo rủi ro được hiển thị bằng màu sắc.",
        ),
    )

    for image, caption_text in illustrations:
        image_paragraph = doc.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_before = Pt(4)
        image_paragraph.paragraph_format.space_after = Pt(1)
        image_paragraph.add_run().add_picture(image, width=Cm(15.2))
        anchor._p.addnext(image_paragraph._p)

        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(5)
        run = caption.add_run(caption_text)
        _set_run_font(run, size=9, color=DARK_GREEN)
        run.italic = True
        image_paragraph._p.addnext(caption._p)
        anchor = caption

    placeholder._element.getparent().remove(placeholder._element)
    if old_caption is not None and old_caption.text.startswith("Hình 1. Lưu đồ vận hành"):
        old_caption._element.getparent().remove(old_caption._element)


def apply_report_design(doc):
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        for run in paragraph.runs:
            _set_run_font(run, size=13)
        text = paragraph.text.strip()
        if text.startswith(("BÁO CÁO KẾT QUẢ", "CUỘC THI KHOA HỌC", "DỰ ÁN: SAFEWEB")):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                _set_run_font(run, size=14, color=DARK_GREEN, bold=True)
        elif text.startswith("BÁO CÁO TÓM TẮT"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            for run in paragraph.runs:
                _set_run_font(run, size=14, color=DARK_GREEN, bold=True)
        elif text.startswith(("1. VẤN ĐỀ", "2. THIẾT KẾ", "3. THỰC HIỆN", "4. KẾT QUẢ",
                              "TÀI LIỆU THAM KHẢO", "PHỤ LỤC:")):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(9)
            paragraph.paragraph_format.space_after = Pt(5)
            for run in paragraph.runs:
                _set_run_font(run, size=13, color=DARK_GREEN, bold=True)

    for table in doc.tables:
        table.autofit = True
        is_model_table = len(table.columns) >= 6
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                _set_cell_margins(cell)
                _set_cell_shading(
                    cell,
                    DARK_GREEN if row_index == 0 else (PALE_GREEN if row_index % 2 else PALE_TEAL),
                )
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        _set_run_font(
                            run,
                            size=8 if is_model_table else 10,
                            color="FFFFFF" if row_index == 0 else TEXT,
                            bold=True if row_index == 0 else None,
                        )
        if table.rows:
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    _insert_illustrations(doc)
