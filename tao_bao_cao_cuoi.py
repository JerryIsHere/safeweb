import os
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from tao_bao_cao_cuoi2 import apply_report_design

def create_report():
    doc = docx.Document()
    
    # 1. Cài đặt chuẩn Phụ lục III: Lề A4 Trái 3cm, Phải 2cm, Trên 2cm, Dưới 2cm
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    # Hàm hỗ trợ ghi văn bản chuẩn: Times New Roman, cỡ 13pt, cách dòng đơn, cách đoạn 3pt
    def add_p(text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.alignment = align
        if text:
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            r.bold = bold
            r.italic = italic
        return p

    def add_table(headers, rows):
        tbl = doc.add_table(rows=len(rows)+1, cols=len(headers))
        tbl.style = 'Table Grid'
        for j, h in enumerate(headers):
            p = tbl.cell(0, j).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            r.bold = True
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                p = tbl.cell(i+1, j).paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(val)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(13)
        add_p("") # khoảng trắng sau bảng

    # ==================== TRANG BÌA ====================
    add_p("BÁO CÁO KẾT QUẢ THỰC HIỆN DỰ ÁN DỰ THI (IIIa)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p("CUỘC THI KHOA HỌC KỸ THUẬT CẤP TRUNG HỌC PHỔ THÔNG", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph() # Cách dòng
    add_p("DỰ ÁN: SAFEWEB - ỨNG DỤNG TRÍ TUỆ NHÂN TẠO KẾT HỢP PHÂN TÍCH ĐỊA CHỈ TRANG WEB GIÚP CẢNH BÁO SỚM RỦI RO LỪA ĐẢO VÀ LỘ LỌT DỮ LIỆU", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    add_p("Lĩnh vực: Máy tính / Công nghệ thông tin", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p("Dạng dự án: Dự án Kỹ thuật", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    # ==================== BÁO CÁO TÓM TẮT ====================
    add_p("BÁO CÁO TÓM TẮT KẾT QUẢ THỰC HIỆN DỰ ÁN DỰ THI (IIIb)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p("Tên dự án dự thi: SAFEWEB - ỨNG DỤNG TRÍ TUỆ NHÂN TẠO KẾT HỢP PHÂN TÍCH ĐỊA CHỈ TRANG WEB GIÚP CẢNH BÁO SỚM RỦI RO LỪA ĐẢO VÀ LỘ LỌT DỮ LIỆU", bold=True)
    add_p("Lĩnh vực dự thi: Máy tính / Công nghệ thông tin", bold=True)
    add_p("Tóm tắt kết quả:", bold=True)
    add_p("- Mục tiêu: Xây dựng công cụ kiểm tra nhanh địa chỉ ngay trên trình duyệt, phát hiện sớm lừa đảo và nêu rõ lý do cảnh báo mà không làm lộ thông tin riêng tư.")
    add_p("- Phương pháp nghiên cứu: Đo 55 dấu hiệu cấu trúc, dùng bốn mô hình học máy (Random Forest, Gradient Boosting, SVM, Logistic Regression) cùng cơ sở dữ liệu 36.261 địa chỉ; thử trên các tên miền không trùng lúc học.")
    add_p("- Quá trình thực hiện: Chuẩn bị 261.299 địa chỉ đã làm sạch để huấn luyện; xây dựng quy trình phân tích trực tiếp trên máy người dùng (đo đạc, đối chiếu dữ liệu, chấm điểm AI) và xuất kết quả cảnh báo.")
    add_p("- Kết quả chính: Trên 6.624 địa chỉ thi thử mới, mô hình Gradient Boosting bắt trúng 95,61% trang lừa đảo (2.676/2.799); phần mềm đạt 100% qua 59 bài kiểm tra xử lý dữ liệu và 16 bài kiểm thử trình duyệt.")
    add_p("- Ý nghĩa/Kết luận: Giúp người dùng chủ động nhận biết rủi ro trực quan kèm lý do minh bạch trước khi bấm vào đường dẫn, hình thành thói quen kiểm tra an toàn số.")
    doc.add_page_break()

    # ==================== NỘI DUNG CHÍNH (ĐẮP TỪ FILE CŨ) ====================
    
    add_p("1. VẤN ĐỀ NGHIÊN CỨU HOẶC CẦN GIẢI QUYẾT", bold=True)
    add_p("Mỗi trang web có một dòng địa chỉ giống như biển hiệu của cửa hàng. Kẻ gian gửi tin nhắn, thư điện tử chứa đường dẫn giả mạo ngân hàng, mạng xã hội hay sàn mua sắm để lừa người dùng nhập mật khẩu và mã xác nhận. Mắt thường khó nhận ra vì địa chỉ giả thường chỉ khác thật vài ký tự, hoặc viết rất dài để che phần đuôi.")
    add_p("Nhiều người tin rằng địa chỉ có khóa https là an toàn, nhưng kẻ lừa đảo cũng tạo được khóa này miễn phí. Nếu gửi địa chỉ lên một máy chủ lạ để kiểm tra, người dùng còn có nguy cơ lộ thông tin riêng. Vì vậy cần một công cụ kiểm tra ngay trên máy của người dùng và nói rõ lý do cảnh báo.")
    add_p("Nhóm cần một công cụ tự điền đủ địa chỉ gõ thiếu, đối chiếu với cơ sở dữ liệu 36.261 địa chỉ đã biết (19.909 trang lừa đảo, 16.352 trang hợp lệ), chấm điểm bằng bốn mô hình và nêu lý do cảnh báo.")
    add_p("Mục tiêu: Xây dựng SafeWeb giúp người dùng kiểm tra nhanh một địa chỉ ngay trên trình duyệt và hiểu vì sao nên hoặc không nên tin trước khi bấm.")
    add_p("Tiêu chí thiết kế: Đo 55 dấu hiệu, dùng bốn mô hình cùng cơ sở dữ liệu 36.261 địa chỉ; thử trên các tên miền không trùng lúc học; mặc định không gửi địa chỉ ra ngoài; hiện bốn mức màu Xanh - Vàng - Cam - Đỏ kèm lý do.")
    add_p("SafeWeb chỉ xem xét chữ viết và cấu trúc của địa chỉ, không tải hay chạy nội dung trang đích. Kiểm tra tên miền (DNS) và ảnh xem trước chỉ chạy khi người dùng bật cho phép.")

    add_p("2. THIẾT KẾ VÀ PHƯƠNG PHÁP NGHIÊN CỨU", bold=True)
    add_p("Thay vì chỉ nghe một ý kiến, SafeWeb dùng bốn mô hình học máy, tức là bốn cách để máy tính học dấu hiệu từ các trang web mẫu. Có thể hình dung chúng như bốn vị giám khảo cùng xem một hồ sơ nhưng mỗi người để ý một kiểu dấu hiệu khác nhau:")
    add_p("- Random Forest: Tạo ra 100 cây câu hỏi nhỏ. Mỗi cây lần lượt hỏi về độ dài, dấu chấm, dấu gạch ngang hay từ khóa. Có thể hình dung 100 giám khảo độc lập bỏ phiếu; kết quả chung dựa trên tiếng nói của cả nhóm, nên một dấu hiệu lạ đơn lẻ khó quyết định toàn bộ kết luận.")
    add_p("- Gradient Boosting: Dùng nhiều cây hỏi nối tiếp nhau theo kiểu người sau xem lại chỗ người trước còn nhầm. Trong phép thử của dự án, mô hình này phát hiện 2.676 trên 2.799 trang lừa đảo, tương ứng 95,61%.")
    add_p("- SVM: Tưởng tượng xếp hồ sơ an toàn về một phía, hồ sơ lừa đảo về phía còn lại rồi tìm đường ranh ngăn cách hai nhóm.")
    add_p("- Logistic Regression: Giống như một bảng cân có nhiều quả cân: mỗi đặc điểm làm thay đổi điểm nghiêng về một phía. Cách tính gọn và nhanh, giúp nhóm có thêm một ý kiến khác để so sánh với ba vị giám khảo còn lại.")
    add_p("Quy trình máy tính biến một dòng địa chỉ trang web thành kết luận cảnh báo trên màn hình được xác định rõ ràng qua 4 bước tính toán cụ thể:")
    add_p("- Bước 1 — Đo đạc thành 55 con số: Phần mềm đếm độ dài, số lượng dấu câu, tỷ lệ chữ số và các từ khóa trên đường dẫn để biến thành 55 con số cụ thể.")
    add_p("- Bước 2 — Tính điểm trung bình của 4 mô hình AI: Từng mô hình dựa vào 55 con số này để tính ra tỷ lệ nghi ngờ lừa đảo từ 0% đến 100%. Hệ thống cộng kết quả của cả 4 mô hình rồi chia đều cho 4 để ra mức xác suất nghi ngờ chung.")
    add_p("- Bước 3 — Cộng điểm các dấu hiệu nhìn thấy được: Hệ thống cộng điểm riêng khi thấy địa chỉ quá dài, nhiều tầng tên phụ, dùng dãy số thay tên, có dấu @, nhiều ký tự lạ, có dấu mã hóa hoặc từ khóa nhạy cảm.")
    add_p("- Bước 4 — Xếp mức rủi ro và đổi màu: Mức xác suất dưới 25% là rủi ro thấp (Xanh); từ 25% đến dưới 50% là cần thận trọng (Vàng); từ 50% đến dưới 75% là nguy cơ cao (Cam); từ 75% trở lên là nguy cơ rất cao (Đỏ).")
    add_p("Những điểm dưới đây là thứ phân biệt SafeWeb với một công cụ chỉ tra danh sách đen hoặc chỉ báo chung chung:")
    add_p("- Ba lớp thông tin cùng làm việc: Hồ sơ đã biết trong cơ sở dữ liệu, 55 dấu hiệu cấu trúc của địa chỉ và ý kiến của bốn mô hình được đặt cạnh nhau.")
    add_p("- Giải thích được vì sao: Người dùng thấy mức màu, điểm của từng mô hình và lý do cụ thể như địa chỉ quá dài, nhiều tầng tên phụ hay có từ khóa nhạy cảm.")
    add_p("- Chạy ngay trong trình duyệt: Việc đối chiếu và chấm điểm diễn ra trên máy người dùng, mặc định không gửi địa chỉ đi đâu.")
    add_p("- Tự sửa địa chỉ gõ thiếu: Nếu người dùng chỉ gõ một từ như vietcombank, phần mềm tự điền thành địa chỉ đầy đủ trước khi kiểm tra.")
    add_p("- Ví dụ lừa đảo tương tự để đối chiếu: Cung cấp trang lừa đảo có tên miền giống để người dùng thấy trang thật khác trang giả thế nào.")
    add_p("- Phiếu báo cáo do người dùng lập: Người dùng có thể lập phiếu báo cáo một địa chỉ nghi lừa đảo. Phiếu được lưu thành tệp trên máy, không tự gửi đi.")
    add_p("Lược đồ 55 dấu hiệu nhận biết địa chỉ trang web được chia thành 6 nhóm: (1) Độ dài và đếm ký tự; (2) Cấu trúc và ký hiệu bất thường; (3) Độ lộn xộn của chữ cái; (4) Từ khóa nhạy cảm; (5) Tỉ lệ chữ, số và đuôi tên miền đáng ngờ; (6) Giấu chữ bằng mã phần trăm.")
    add_p("Bộ đo được viết hai lần: một bản trong Python để dùng khi huấn luyện và một bản trong JavaScript để dùng trên trình duyệt.")

    add_p("3. THỰC HIỆN, CHẾ TẠO VÀ KIỂM TRA", bold=True)
    add_p("Thành phần và công cụ sử dụng: Python, pandas và scikit-learn để xử lý và huấn luyện; HTML, CSS, JavaScript và Flask để làm giao diện và máy chủ.")
    add_p("Dữ liệu đưa vào hệ thống: Nhóm chuẩn bị ba phần dữ liệu. Thứ nhất, cơ sở dữ liệu tra cứu nhanh giữ 36.261 địa chỉ (19.909 lừa đảo và 16.352 hợp lệ). Thứ hai, tập dữ liệu dùng để dạy máy học tổng cộng có 261.299 địa chỉ. Thứ ba, tập dữ liệu giữ riêng để thi thử gồm 6.624 địa chỉ được để riêng, không dùng lúc huấn luyện.")
    add_p("Quy trình huấn luyện máy học: Bước 1: Làm sạch dữ liệu và tách 6.624 địa chỉ thi thử. Bước 2: Đồng bộ nhãn dữ liệu. Bước 3: Đo đạc 55 con số đặc điểm. Bước 4: Cho bốn mô hình học quy luật. Bước 5: Chấm điểm và đóng gói lên trình duyệt.")
    add_p("Quá trình thử nghiệm: Ở giai đoạn đầu, nhóm dùng khoảng 5.000 địa chỉ và một mô hình Random Forest. Ở các lần nâng cấp, nhóm mở rộng lên 261.299 địa chỉ và dùng bốn mô hình.")
    add_p("Kết quả định lượng của bốn mô hình Trí tuệ nhân tạo trên 6.624 địa chỉ thi thử:")
    
    add_table(
        ["Tên mô hình", "Tỷ lệ đoán đúng chung", "Độ chuẩn xác khi báo động", "Tỷ lệ bắt trúng lừa đảo", "Điểm hiệu quả tổng hợp", "Đúng / Nhầm / Sót / Bắt trúng"],
        [
            ["1. Random Forest", "87,45%", "80,75%", "92,32%", "86,15%", "3.209 / 616 / 215 / 2.584"],
            ["2. Gradient Boosting", "87,89%", "79,76%", "95,61%", "86,97%", "3.146 / 679 / 123 / 2.676"],
            ["3. SVM", "58,05%", "100,00%", "0,71%", "1,42%", "3.825 / 0 / 2.779 / 20"],
            ["4. Logistic Regression", "85,55%", "79,96%", "87,82%", "83,71%", "3.209 / 616 / 341 / 2.458"]
        ]
    )
    
    add_p("Kiểm thử phần mềm và xác minh cục bộ: Đạt 59/59 bài kiểm thử xử lý dữ liệu trên máy tính (100%); đạt 16/16 bài kiểm thử trên trình duyệt web. Việc đối chiếu hai cách tính giữa Python và JavaScript khớp tuyệt đối về mặt phân loại. Xác minh cục bộ đảm bảo địa chỉ được xử lý ngay trong trình duyệt, không tự gửi đi dịch vụ bên ngoài.")
    
    add_p("[BẠN CÓ THỂ CHÈN ẢNH LƯU ĐỒ, ẢNH XÁC MINH VÀ ẢNH GIAO DIỆN VÀO ĐÂY]", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p("Hình 1. Lưu đồ vận hành và kết quả xác minh cục bộ (Đồ họa do học sinh tự thực hiện).", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)

    add_p("Hướng dẫn thao tác: Người dùng thực hiện 3 bước: Dán địa chỉ vào ô nhập; Bấm nút phân tích URL; Nhìn màu sắc cảnh báo và đọc lý do cụ thể.")

    add_p("4. KẾT QUẢ, KẾT LUẬN VÀ KHẢ NĂNG ỨNG DỤNG", bold=True)
    add_p("Khẳng định đóng góp mới và ưu thế vượt trội: SafeWeb đặt cạnh nhau ba loại thông tin (hồ sơ, dấu hiệu cấu trúc, ý kiến bốn mô hình). Giúp người dùng xem xét trước khi bấm nhờ màu cảnh báo, điểm từng mô hình và lý do minh bạch.")
    add_p("Kết luận và một số hạn chế: Trong phép thử, mô hình Gradient Boosting phát hiện 95,61% trang lừa đảo. SVM còn bỏ sót nhiều trường hợp và cần được hiệu chỉnh thêm.")
    add_p("Hướng phát triển tương lai: Phát triển SafeWeb thành tiện ích gắn trực tiếp vào trình duyệt để tự động rà các đường dẫn đang hiện trên màn hình và đặt dấu cảnh báo trước khi người dùng bấm.")
    add_p("Giá trị và định hướng giáo dục: SafeWeb còn là bài học an toàn số giúp người dùng rèn luyện thói quen kiểm tra trước khi bấm.")

    # ==================== TÀI LIỆU THAM KHẢO & PHỤ LỤC ====================
    add_p("TÀI LIỆU THAM KHẢO", bold=True)
    add_p("1. Bộ dữ liệu địa chỉ trang web tiếng Việt đã kiểm chứng: dùng làm cơ sở dữ liệu tra cứu nhanh và bộ đề thi đánh giá độc lập.")
    add_p("2. Bộ dữ liệu địa chỉ trang web lừa đảo mở rộng quốc tế: dùng để bổ sung nhiều kiểu địa chỉ giúp các mô hình học thêm dấu hiệu nhận diện.")
    add_p("3. Tài liệu kỹ thuật thư viện Học máy scikit-learn (Ngôn ngữ Python): dùng để huấn luyện, chấm điểm và thống kê kết quả bốn mô hình Trí tuệ nhân tạo.")
    add_p("4. Chuẩn kỹ thuật Quốc tế về cấu trúc địa chỉ web (RFC 3986) và hệ thống tên miền (RFC 1035): dùng để tách và chuẩn hóa các thành phần của địa chỉ trang web.")

    add_p("PHỤ LỤC: TÁM TÊN GỌI CỐT LÕI", bold=True)
    add_p("1. URL: Địa chỉ để đi đến một trang, giống số nhà và tên đường.")
    add_p("2. Phishing: Trang giả mạo nhằm dụ người xem giao mật khẩu hoặc tiền.")
    add_p("3. Domain: Tên dễ nhớ thay cho dãy số của máy chủ.")
    add_p("4. HTTPS: Lớp mã hóa đường truyền; có khóa không đồng nghĩa trang là thật.")
    add_p("5. DNS: Dịch vụ đổi tên trang thành địa chỉ máy chủ.")
    add_p("6. Machine Learning: Cho máy xem nhiều ví dụ để nó học cách nhận ra mẫu.")
    add_p("7. Browser inference: Máy người dùng tự chấm điểm, như tự làm phép tính tại bàn.")
    add_p("8. Browser extension: Chương trình nhỏ thêm chức năng vào trình duyệt.")

    apply_report_design(doc)
    doc.save("Bao_cao_SafeWeb_KHKT_Chuan_Moi.docx")
    print("Đã tạo Bao_cao_SafeWeb_KHKT_Chuan_Moi.docx với nội dung báo cáo và định dạng SafeWeb.")

if __name__ == '__main__':
    create_report()