"""Generate the report directly from the user's supplied report text."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports" / "approved_report_source.txt"
OUTPUT = ROOT / "Bao_cao_SafeWeb_KHKT_THPT.docx"
FIGURES = ROOT / "reports" / "figures_vi"
FONT = "Times New Roman"


def text_blocks(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"\s{2,}", text.strip()) if part.strip()]


def section(text: str, start: str, end: str | None = None) -> str:
    left = text.index(start)
    right = text.index(end, left + len(start)) if end else len(text)
    return text[left:right]


def add_run(paragraph, text: str, *, bold: bool = False, italic: bool = False, size: int = 13) -> None:
    run = paragraph.add_run(text)
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_paragraph(document, text: str, *, align=None, bold: bool = False, size: int = 13, italic: bool = False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_together = True
    if align is not None:
        paragraph.alignment = align
    add_run(paragraph, text, bold=bold, italic=italic, size=size)
    return paragraph


def add_lead_paragraph(document, text: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_together = True
    match = re.match(r"(.{1,100}?:)(\s*)(.*)", text, flags=re.DOTALL)
    if match:
        add_run(paragraph, match.group(1), bold=True)
        add_run(paragraph, match.group(2) + match.group(3))
    else:
        add_run(paragraph, text)
    return paragraph


def add_heading(document, text: str, *, major: bool = False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(6 if major else 3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_run(paragraph, text, bold=True, size=14)
    return paragraph


def extract_prefixed(block: str, marker: str) -> tuple[str, str]:
    if not block.startswith(marker):
        raise ValueError(f"Expected section text to start with {marker!r}")
    return marker, block[len(marker):].strip()


def make_document(text: str) -> None:
    document = Document()
    page = document.sections[0]
    page.page_width = Cm(21)
    page.page_height = Cm(29.7)
    page.left_margin = Cm(3)
    page.right_margin = Cm(2)
    page.top_margin = Cm(2)
    page.bottom_margin = Cm(2)
    for style in document.styles:
        if hasattr(style, "font"):
            style.font.name = FONT
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
            style.font.size = Pt(13)

    cover = section(text, "BÁO CÁO KẾT QUẢ THỰC HIỆN DỰ ÁN DỰ THI (IIIa)", "BÁO CÁO TÓM TẮT")
    cover_lines = (
        "BÁO CÁO KẾT QUẢ THỰC HIỆN DỰ ÁN DỰ THI (IIIa)",
        "CUỘC THI KHOA HỌC KỸ THUẬT CẤP TRUNG HỌC PHỔ THÔNG",
        "DỰ ÁN: SAFEWEB - ỨNG DỤNG TRÍ TUỆ NHÂN TẠO KẾT HỢP PHÂN TÍCH ĐỊA CHỈ TRANG WEB GIÚP CẢNH BÁO SỚM RỦI RO LỪA ĐẢO VÀ LỘ LỌT DỮ LIỆU",
        "Lĩnh vực dự thi: Máy tính / Công nghệ thông tin",
        "Dạng dự án: Dự án Kỹ thuật",
    )
    if not all(line in cover for line in cover_lines):
        raise ValueError("The supplied cover text is incomplete.")
    for line in cover_lines:
        add_paragraph(document, line, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14)
    document.add_page_break()

    summary = section(text, "BÁO CÁO TÓM TẮT", "1. VẤN ĐỀ NGHIÊN CỨU")
    summary_headings = (
        "BÁO CÁO TÓM TẮT KẾT QUẢ THỰC HIỆN DỰ ÁN DỰ THI (IIIb)",
        "Tên dự án dự thi:",
        "Lĩnh vực dự thi:",
        "Tóm tắt kết quả:",
    )
    for heading in summary_headings:
        if heading not in summary:
            raise ValueError(f"Missing summary content: {heading}")
    add_heading(document, summary_headings[0], major=True)
    summary_tail = summary.split(summary_headings[0], 1)[1]
    summary_tail = summary_tail.replace("Tên dự án dự thi:", "\nTên dự án dự thi:", 1)
    summary_tail = summary_tail.replace("Lĩnh vực dự thi:", "\nLĩnh vực dự thi:", 1)
    summary_tail = summary_tail.replace("Tóm tắt kết quả:", "\nTóm tắt kết quả:", 1)
    for block in [part.strip() for part in summary_tail.splitlines() if part.strip()]:
        if block == "Tóm tắt kết quả:":
            add_heading(document, block)
        elif block.startswith(("Tên dự án dự thi:", "Lĩnh vực dự thi:")):
            add_lead_paragraph(document, block)
        else:
            for item in re.split(r"(?=(?:Mục tiêu|Phương pháp nghiên cứu|Quá trình thực hiện|Kết quả chính|Ý nghĩa/Kết luận):)", block):
                item = item.strip()
                if item:
                    add_lead_paragraph(document, item)
    document.add_page_break()

    first = section(text, "1. VẤN ĐỀ NGHIÊN CỨU", "2. THIẾT KẾ VÀ PHƯƠNG PHÁP")
    first_title = "1. VẤN ĐỀ NGHIÊN CỨU HOẶC CẦN GIẢI QUYẾT"
    add_heading(document, first_title, major=True)
    first_body = first[len(first_title):]
    labels = (
        "Đòi hỏi của thực tế và sự cấp thiết của vấn đề cần giải quyết:",
        "Xác định các tiêu chí cho giải pháp giải quyết vấn đề:",
    )
    for index, block in enumerate(text_blocks(first_body)):
        if index == 0 and block.startswith(labels[0]):
            add_lead_paragraph(document, block)
        elif block.startswith(labels[1]):
            add_lead_paragraph(document, block)
        else:
            add_paragraph(document, block)

    design = section(text, "2. THIẾT KẾ VÀ PHƯƠNG PHÁP", "3. THỰC HIỆN, CHẾ TẠO")
    design_title = "2. THIẾT KẾ VÀ PHƯƠNG PHÁP NGHIÊN CỨU"
    add_heading(document, design_title, major=True)
    design_body = design[len(design_title):]
    first_design_lead = "Quá trình nghiên cứu và tìm tòi giải pháp (Phối hợp 4 mô hình AI):"
    prototype_lead = "Thiết kế mô hình/nguyên mẫu và quy trình đưa ra phán đoán:"
    for index, block in enumerate(text_blocks(design_body)):
        if index == 0 and block.startswith(first_design_lead):
            add_lead_paragraph(document, block)
        elif block.startswith(prototype_lead):
            add_lead_paragraph(document, block)
        else:
            add_lead_paragraph(document, block)

    implementation = section(text, "3. THỰC HIỆN, CHẾ TẠO", "4. KẾT QUẢ, KẾT LUẬN")
    implementation_title = "3. THỰC HIỆN, CHẾ TẠO VÀ KIỂM TRA"
    add_heading(document, implementation_title, major=True)
    implementation_body = implementation[len(implementation_title):]
    for block in text_blocks(implementation_body):
        if block.startswith("Thực hiện chế tạo và quy trình công nghệ:"):
            add_lead_paragraph(document, block)
            continue
        if block.startswith("Cơ sở dữ liệu tra cứu nhanh:"):
            block = re.sub(r"(?<=\))(?=Tập dữ liệu)|(?<=mô hình\.)(?=Tập dữ liệu)|(?<=trình duyệt web\.)(?=Quy trình)", "\n", block)
            for item in block.splitlines():
                if item.startswith("Quy trình huấn luyện máy học diễn ra"):
                    add_lead_paragraph(document, item)
                else:
                    add_lead_paragraph(document, item)
            continue
        if block.startswith("Kiểm tra và hoàn thiện sản phẩm:"):
            add_lead_paragraph(document, block)
            continue
        if block.startswith(("Random Forest:", "Gradient Boosting:", "SVM:", "Logistic Regression:")):
            for item in re.split(r"(?=(?:Random Forest|Gradient Boosting|SVM|Logistic Regression):)", block):
                if item.strip():
                    add_lead_paragraph(document, item.strip())
            continue
        if block.startswith("Hướng dẫn thao tác:"):
            add_lead_paragraph(document, block)
            continue
        if block.startswith("(Hình 1."):
            image_path = FIGURES / "hinh_3_1_luu_do_vi.png"
            if image_path.is_file():
                picture = document.add_paragraph()
                picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture.paragraph_format.space_after = Pt(0)
                picture.add_run().add_picture(str(image_path), width=Cm(15.5))
            add_paragraph(document, block, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
            continue
        if block.startswith("(Hình 2."):
            for image_name, width in (("hinh_4_3_giao_dien_vi.png", 15), ("hinh_4_2_ba_buoc_vi.png", 15)):
                image_path = FIGURES / image_name
                if image_path.is_file():
                    picture = document.add_paragraph()
                    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    picture.paragraph_format.space_after = Pt(0)
                    picture.add_run().add_picture(str(image_path), width=Cm(width))
            add_paragraph(document, block, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
            continue
        add_paragraph(document, block)

    conclusions = section(text, "4. KẾT QUẢ, KẾT LUẬN", "TÀI LIỆU THAM KHẢO")
    conclusion_title = "4. KẾT QUẢ, KẾT LUẬN VÀ KHẢ NĂNG ỨNG DỤNG"
    add_heading(document, conclusion_title, major=True)
    conclusion_body = conclusions[len(conclusion_title):]
    for block in text_blocks(conclusion_body):
        add_lead_paragraph(document, block)

    references = section(text, "TÀI LIỆU THAM KHẢO", "PHỤ LỤC:")
    add_heading(document, "TÀI LIỆU THAM KHẢO", major=True)
    for block in text_blocks(references[len("TÀI LIỆU THAM KHẢO"):]):
        add_paragraph(document, block)

    glossary = section(text, "PHỤ LỤC:", None)
    glossary_title = "PHỤ LỤC: TÁM TÊN GỌI CỐT LÕI"
    add_heading(document, glossary_title, major=True)
    glossary_body = glossary[len(glossary_title):]
    for block in text_blocks(glossary_body):
        add_lead_paragraph(document, block)

    properties = document.core_properties
    properties.title = "Báo cáo kết quả thực hiện dự án dự thi SafeWeb"
    properties.subject = "Báo cáo dự án kỹ thuật"
    properties.author = ""
    properties.last_modified_by = ""
    properties.keywords = "SafeWeb; dự án kỹ thuật; an toàn số"

    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if "PhishVN" in full_text or "PhiUSIIL" in full_text or re.search(r"https?://|www\.", full_text, re.I):
        raise ValueError("Sensitive dataset names or external URLs remain in the generated report.")
    document.save(OUTPUT)
    print(f"Generated {OUTPUT} ({OUTPUT.stat().st_size:,} bytes).")


if __name__ == "__main__":
    make_document(SOURCE.read_text(encoding="utf-8"))
